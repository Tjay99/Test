from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


DEFAULT_OUTPUT_DIR = Path("output_cvs")


def add_header(doc: Document, name: str, title: str, contact: str) -> None:
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = header.add_run(name)
    name_run.bold = True
    name_run.font.size = Pt(20)
    header.add_run(f"\n{title}")
    header.add_run(f"\n{contact}")


def add_section_title(doc: Document, title: str) -> None:
    section = doc.add_paragraph()
    run = section.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(12)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_experience(doc: Document, entries: list[dict[str, object]]) -> None:
    for entry in entries:
        role = entry["role"]
        company = entry["company"]
        dates = entry["dates"]
        summary = entry["summary"]
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"{role} — {company}")
        run.bold = True
        paragraph.add_run(f"\n{dates}")
        if summary:
            add_bullets(doc, summary)


def add_education(doc: Document, entries: list[dict[str, str]]) -> None:
    for entry in entries:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"{entry['degree']} — {entry['school']}")
        run.bold = True
        paragraph.add_run(f"\n{entry['dates']}")


def build_modern_template(output_path: Path) -> None:
    doc = Document()
    add_header(
        doc,
        "Camille Dupont",
        "Chef de projet digital",
        "camille.dupont@email.com | +33 6 12 34 56 78 | Paris",
    )
    doc.add_paragraph(
        "Professionnel(le) du digital avec 6 ans d'expérience en conduite de projets "
        "web et mobile. Spécialisé(e) dans la coordination d'équipes pluridisciplinaires "
        "et l'amélioration continue."
    )

    add_section_title(doc, "Expérience")
    add_experience(
        doc,
        [
            {
                "role": "Chef de projet digital",
                "company": "Agence Nova",
                "dates": "2021 - Aujourd'hui",
                "summary": [
                    "Pilotage de 12 projets web avec un taux de livraison à 96%.",
                    "Mise en place d'un tableau de bord KPI pour 5 clients majeurs.",
                ],
            },
            {
                "role": "Chargé(e) de projet",
                "company": "Start-up Lumen",
                "dates": "2018 - 2021",
                "summary": [
                    "Coordination de l'équipe produit (UX, dev, marketing).",
                    "Optimisation du cycle de développement, -20% sur les délais.",
                ],
            },
        ],
    )

    add_section_title(doc, "Compétences")
    add_bullets(
        doc,
        [
            "Gestion Agile (Scrum, Kanban)",
            "Planification budgétaire",
            "Outils: Jira, Notion, Figma",
            "Communication client et reporting",
        ],
    )

    add_section_title(doc, "Formation")
    add_education(
        doc,
        [
            {
                "degree": "Master Management Digital",
                "school": "Université Paris Dauphine",
                "dates": "2016 - 2018",
            }
        ],
    )

    doc.save(output_path)


def build_creative_template(output_path: Path) -> None:
    doc = Document()
    add_header(
        doc,
        "Nora Benali",
        "Designer UI/UX",
        "nora.benali@email.com | +33 7 98 76 54 32 | Lyon",
    )
    doc.add_paragraph(
        "Designer UI/UX passionné(e) par la création d'expériences utilisateur inclusives "
        "et l'élaboration d'identités visuelles modernes."
    )

    add_section_title(doc, "Projets clés")
    add_experience(
        doc,
        [
            {
                "role": "Lead UI/UX Designer",
                "company": "Studio Vivid",
                "dates": "2022 - Aujourd'hui",
                "summary": [
                    "Refonte complète d'une application mobile B2C (4,8/5 sur les stores).",
                    "Création d'un design system pour 3 produits SaaS.",
                ],
            },
            {
                "role": "Designer Produit",
                "company": "Maison Oria",
                "dates": "2019 - 2022",
                "summary": [
                    "Animation d'ateliers de co-création avec les équipes métier.",
                    "Prototype interactif livré en 3 semaines pour un MVP retail.",
                ],
            },
        ],
    )

    add_section_title(doc, "Compétences")
    add_bullets(
        doc,
        [
            "Design system & wireframing",
            "Recherche utilisateur qualitative",
            "Outils: Figma, Adobe XD, Miro",
            "Accessibilité numérique",
        ],
    )

    add_section_title(doc, "Formation")
    add_education(
        doc,
        [
            {
                "degree": "Diplôme Supérieur d'Arts Appliqués",
                "school": "Lycée La Martinière",
                "dates": "2015 - 2018",
            }
        ],
    )

    doc.save(output_path)


def build_executive_template(output_path: Path) -> None:
    doc = Document()
    add_header(
        doc,
        "Marc Lefevre",
        "Directeur Commercial",
        "marc.lefevre@email.com | +33 6 55 44 33 22 | Bordeaux",
    )
    doc.add_paragraph(
        "Leader commercial avec 12 ans d'expérience dans la croissance de revenus "
        "et le management d'équipes internationales."
    )

    add_section_title(doc, "Expérience")
    add_experience(
        doc,
        [
            {
                "role": "Directeur Commercial",
                "company": "Groupe Horizon",
                "dates": "2019 - Aujourd'hui",
                "summary": [
                    "Augmentation du chiffre d'affaires de 35% en 24 mois.",
                    "Structuration d'une équipe de 18 commerciaux sur 4 régions.",
                ],
            },
            {
                "role": "Responsable des ventes",
                "company": "TechLine",
                "dates": "2014 - 2019",
                "summary": [
                    "Déploiement d'une stratégie de prospection multicanal.",
                    "Formation et coaching des équipes terrain.",
                ],
            },
        ],
    )

    add_section_title(doc, "Compétences")
    add_bullets(
        doc,
        [
            "Stratégie commerciale B2B",
            "Management d'équipes multi-sites",
            "Négociation grands comptes",
            "Analyse de performance",
        ],
    )

    add_section_title(doc, "Formation")
    add_education(
        doc,
        [
            {
                "degree": "MBA Commerce International",
                "school": "KEDGE Business School",
                "dates": "2012 - 2014",
            }
        ],
    )

    doc.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Génère 3 modèles de CV modernes au format Word (.docx)."
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Dossier de sortie pour les fichiers .docx.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_dir: Path = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    build_modern_template(output_dir / "cv_moderne_chef_projet.docx")
    build_creative_template(output_dir / "cv_creatif_designer.docx")
    build_executive_template(output_dir / "cv_executive_directeur.docx")

    print(f"3 modèles créés dans: {output_dir.resolve()}")


if __name__ == "__main__":
    main()
